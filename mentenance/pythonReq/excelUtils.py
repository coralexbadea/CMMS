from openpyxl import *
from openpyxl.utils.cell import *
import xlsxwriter
import pandas as pd

class ExcelHandler:
	def __init__(self, excel_path="./test.xlsx", sheet="Sheet1"):
		self.excel_path = excel_path
		self.wb = load_workbook(excel_path)
		self.sheet = self.wb[sheet]

	def read_excel_table(self, table_name="Table1", isTable=True, table_range="A1:A2"):
		if(isTable):
			table = self.sheet.tables[table_name]
			self.table_range = table.ref
		else:
			self.table_range = table_range

		table_head = self.sheet[table_range][0]
		table_data = self.sheet[table_range][1:]

		columns = [column.value for column in table_head]
		data = {column: [] for column in columns}

		for row in table_data:
			row_val = [cell.value for cell in row]
			for key, val in zip(columns, row_val):
				data[key].append(str(val))

		return columns, data, table_range

	def setSheet(self, sheet):
		self.sheet = sheet

	def parseRef(self, refs):   
		letter1, number1 = coordinate_from_string(refs.split(":")[0])
		letter2, number2 = coordinate_from_string(refs.split(":")[1])
		return letter1, number1, letter2, number2

	def insert_row(self, data, letter, number):
		for i, entry in enumerate(data):
			self.sheet.cell(row=number, column=i+column_index_from_string(letter), value=entry)
	
	def extend_table(self, l1, n1, l2, n2):
		self.table_range = "".join([l1,str(n1),":",l2,str(n2)])
		print(self.table_range)


	def save(self):
		self.wb.save(self.excel_path)


from os import path
from glob import glob  
def find_ext(dr, ext):
	return glob(path.join(dr,"*.{}".format(ext)))
	#return [path.basename(path.normpath(p)) for p in paths]